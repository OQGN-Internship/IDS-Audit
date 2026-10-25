"""The Excel exports (ENG-03 9), each read back with openpyxl and compared
with the tables its screen builds, cell by cell, with the Source sheet,
the cells' types, the row limit, the file name, the permission check and
the audit event.
"""

import datetime
import io
import re
from decimal import Decimal

import openpyxl
import pytest
from django.test import Client
from django.urls import reverse

from ids_audit.accounts import access
from ids_audit.audit.models import AuditEvent
from ids_audit.imports.models import Upload
from ids_audit.ledger.models import GlCiMap, Posting
from ids_audit.reports import export, faq, screens
from ids_audit.reports.models import Report
from ids_audit.reports.queries import ALL, Filters, Keys
from tests.factories import add_posting, commit, make_report, make_upload, make_user
from tests.test_queries import JUL

pytestmark = pytest.mark.django_db


@pytest.fixture
def admin():
    return make_user()


@pytest.fixture
def client(admin) -> Client:
    signed_in = Client()
    signed_in.force_login(admin)
    return signed_in


def download(client: Client, name: str, report: Report, *args, query: str = ""):
    """Export a screen and return the response and its workbook."""
    response = client.get(reverse(f"{name}_export", args=[report.pk, *args]) + query)
    assert response.status_code == 200, response.status_code
    data = b"".join(response.streaming_content)
    return response, openpyxl.load_workbook(io.BytesIO(data))


def read_back(cell: screens.Cell):
    """What openpyxl reads back for a screen's cell (ENG-03 9.3)."""
    value = cell.value
    if value is None or value == "":
        return None
    if cell.kind in ("amount", "exact", "share"):
        return float(Decimal(value))
    if cell.kind == "count":
        return int(value)
    if cell.kind == "date":
        return datetime.datetime.combine(value, datetime.time())
    return str(value)


def sheet_rows(sheet) -> list[list]:
    """A sheet's rows as values, trailing empty cells dropped."""
    found = []
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        while values and values[-1] is None:
            values.pop()
        found.append(values)
    return found


def expected_rows(table: screens.Table) -> list[list]:
    """A table as its sheet should hold it: the header, every row, the total."""
    rows = [[column.title for column in table.columns]]
    rows += [[read_back(cell) for cell in row] for row in table.rows]
    if table.total is not None:
        rows.append([read_back(cell) for cell in table.total])
    for row in rows:
        while row and row[-1] is None:
            row.pop()
    return rows


SCREENS = [
    ("overview", [], lambda r: screens.overview(r, ALL)),
    ("gl_accounts", [], lambda r: screens.gl_accounts(r, ALL)),
    ("gl_account", ["610100"], lambda r: screens.gl_account(r, ALL, "610100")),
    ("departments", [], lambda r: screens.departments(r, ALL)),
    ("department", ["outside"], lambda r: screens.department(r, ALL, "outside")),
    ("budget_lines", [], lambda r: screens.budget_lines(r, ALL)),
    ("budget_line", ["none"], lambda r: screens.budget_line(r, ALL, "none")),
    ("rows", [], lambda r: screens.source_rows(r, ALL, Keys(), None)),
    ("exceptions", [], lambda r: screens.exceptions(r, ALL)),
]


@pytest.mark.parametrize(("name", "args", "build"), SCREENS)
def test_every_export_has_its_screen_s_rows_order_and_totals(
    client, sample_report, name, args, build
) -> None:
    """Verify each screen's export holds one sheet per table, in order, then
    Source, and each sheet the table's header, rows and total row, cell by
    cell (ENG-03 9).
    """
    _, workbook = download(client, name, sample_report, *args)
    screen = build(sample_report)
    assert workbook.sheetnames == [t.sheet for t in screen.tables] + ["Source"]
    for table in screen.tables:
        sheet = workbook[table.sheet]
        assert sheet_rows(sheet) == expected_rows(table), table.sheet
        assert sheet.freeze_panes == "A2"


@pytest.mark.parametrize("entry", faq.ENTRIES, ids=lambda e: e.slug)
def test_every_faq_page_exports_its_tables_question_and_answer(
    client, sample_report, entry
) -> None:
    """Verify each FAQ page's export holds its tables, and the Source sheet
    its question and answer (ENG-03 7 and 9.2).
    """
    _, workbook = download(client, "faq_page", sample_report, entry.slug)
    page = faq.page(sample_report, Filters(), entry, None)
    for table in page.tables:
        assert sheet_rows(workbook[table.sheet]) == expected_rows(table)
    source = dict(
        row[:2] for row in workbook["Source"].iter_rows(values_only=True) if row[0]
    )
    assert source["Question"] == entry.question
    assert source["Answer"] == page.answer


def source_facts(workbook) -> dict:
    """The Source sheet's facts, by label, up to its uploads."""
    found = {}
    for row in workbook["Source"].iter_rows(values_only=True):
        if row[0] == "Committed uploads":
            break
        if row[0]:
            found[row[0]] = row[1]
    return found


def test_the_source_sheet_names_the_filters_uploads_user_and_version(
    client, admin, sample_report, sample_expected, settings
) -> None:
    """Verify the Source sheet states the report, view, filters, as-of line,
    who exported it as Full name (username), and the app's version tag, and
    lists every committed upload with what it counts under the filters,
    adding up to the total (ENG-03 9.2).
    """
    settings.APP_VERSION = "v1.2.3"
    _, workbook = download(
        client, "gl_account", sample_report, "610100", query="?year=2025"
    )
    facts = source_facts(workbook)
    assert facts["Report"] == "Sample"
    assert facts["View"] == "GL account 610100: Software licences"
    assert facts["Address"] == f"/reports/{sample_report.pk}/gl/610100?year=2025"
    assert facts["Filters"] == "Year: 2025. Outside IDS: included."
    assert facts["Figures as of"].startswith("Upload 1, sample.xlsx, approved by ")
    assert facts["Exported by"] == f"{admin.full_name} ({admin.username})"
    assert facts["App version"] == "v1.2.3"
    rows = sheet_rows(workbook["Source"])
    header = rows.index(
        [
            "Upload",
            "File",
            "SHA-256",
            "Rows read",
            "Approved by",
            "Approved",
            "Ruleset version",
            "Lines counted",
            "Actual spend counted",
        ]
    )
    upload, total = rows[header + 1], rows[header + 2]
    assert upload[:4] == [1, "sample.xlsx", sample_report.uploads.get().file_sha256, 35]
    assert upload[6] == 1
    year_2025 = float(Decimal(sample_expected["screens"]["by_year"]["2025"]))
    assert upload[7:] == [total[7], year_2025] and total[0] == "Total"
    assert total[8] == year_2025


def test_a_replaced_upload_is_listed_with_nothing_counted(client) -> None:
    """Verify an upload whose months a later file replaced is still listed
    on the Source sheet, with 0 lines (ENG-03 9.2, G14).
    """
    report = make_report()
    first, second = make_upload(report), make_upload(report)
    add_posting(first, JUL, "1.000")
    add_posting(second, JUL, "2.000")
    commit(first, minute=1)
    commit(second, minute=2)
    _, workbook = download(client, "overview", report)
    rows = sheet_rows(workbook["Source"])
    listed = [row for row in rows if row and row[0] in (1, 2)]
    assert [(row[0], row[7], row[8]) for row in listed] == [(1, 0, 0.0), (2, 1, 2.0)]


def test_a_formula_and_a_long_number_stay_text(client) -> None:
    """Verify a name holding =1+1 opens as text, not a formula, and a
    10-digit document number as text, unchanged (ENG-03 9.3 and 9.4).
    """
    report = make_report()
    upload = make_upload(report)
    GlCiMap.objects.create(
        upload=upload, report=report, gl_code="610100", gl_name="=1+1", ci_code="2101"
    )
    posting = add_posting(upload, JUL, "1.000")
    Posting.objects.filter(pk=posting.pk).update(
        document_number="5100000123", item_text="https://example.test/invoice"
    )
    commit(upload, minute=1)
    _, workbook = download(client, "gl_accounts", report)
    name = workbook["GL accounts"]["B2"]
    assert (name.value, name.data_type) == ("=1+1", "s")
    _, workbook = download(client, "rows", report)
    number = workbook["Source rows"]["G2"]
    assert (number.value, number.data_type) == ("5100000123", "s")
    link = workbook["Source rows"]["F2"]
    assert (link.value, link.hyperlink) == ("https://example.test/invoice", None)


def test_cells_have_their_formats(client, sample_report) -> None:
    """Verify amounts, shares, dates and totals carry the formats ENG-03 9.3
    gives them.
    """
    _, workbook = download(client, "gl_accounts", sample_report)
    sheet = workbook["GL accounts"]
    assert sheet["E2"].number_format == "#,##0.000;-#,##0.000"
    assert sheet["F2"].number_format == "0.0%"
    total = sheet.cell(row=sheet.max_row, column=5)
    assert total.font.bold and total.number_format == "#,##0.000;-#,##0.000"
    _, workbook = download(client, "rows", sample_report)
    assert workbook["Source rows"]["A2"].number_format == "dd.mm.yyyy"


def test_each_export_writes_one_event_with_its_view_filters_and_rows(
    client, admin, sample_report
) -> None:
    """Verify an export writes exactly one export_downloaded event, keeping
    the view, the filters with its keys, the data rows and the file name
    (ENG-01 6.3).
    """
    response, _ = download(
        client, "rows", sample_report, query="?gl=610100&exclude_outside=1"
    )
    (event,) = AuditEvent.objects.filter(action="export_downloaded")
    assert (event.user, event.report) == (admin, sample_report)
    name = re.search(r'filename="([^"]+)"', response["Content-Disposition"])[1]
    assert event.details == {
        "view": "rows",
        "filters": {"year": "all", "exclude_outside": "1", "gl": "610100"},
        "rows": 4,
        "file_name": name,
    }


def test_the_file_name_is_the_report_view_and_time(client, sample_report) -> None:
    """Verify the file is named <report>_<view>_<yyyymmdd-hhmm>.xlsx, with
    only letters, digits, hyphens and underscores (ENG-03 9.4).
    """
    response, _ = download(client, "department", sample_report, "outside")
    assert response["Content-Disposition"].startswith("attachment;")
    name = re.search(r'filename="([^"]+)"', response["Content-Disposition"])[1]
    assert re.fullmatch(r"Sample_departments-outside_\d{8}-\d{4}\.xlsx", name)
    when = datetime.datetime(2026, 4, 3, 10, 5, tzinfo=datetime.UTC)
    arabic = Report(pk=7, name="تقرير 2026")
    assert export.file_name(arabic, "faq/total-spend", when) == (
        "2026_faq-total-spend_20260403-1405.xlsx"
    )
    assert export.file_name(Report(pk=7, name="تقرير"), "overview", when) == (
        "report7_overview_20260403-1405.xlsx"
    )


def test_an_export_past_the_row_limit_is_refused_with_no_event(
    client, sample_report, monkeypatch
) -> None:
    """Verify an export with more rows than the limit is refused with its
    message, sends no file and writes no event (ENG-03 9.4). The limit is
    lowered, since a million rows would take minutes.
    """
    monkeypatch.setattr(export, "MAX_DATA_ROWS", 34)
    response = client.get(reverse("rows_export", args=[sample_report.pk]))
    assert response.status_code == 200
    assert "Content-Disposition" not in response
    text = response.content.decode()
    assert "This export has 35 rows, more than Excel" in text
    assert not AuditEvent.objects.filter(action="export_downloaded").exists()
    monkeypatch.setattr(export, "MAX_DATA_ROWS", 35)
    download(client, "rows", sample_report)


def test_an_export_of_an_empty_report_has_its_sheets_and_says_why(client) -> None:
    """Verify an export of a report with no committed upload still has every
    sheet, with header and total rows only, and the Source sheet says why.
    """
    _, workbook = download(client, "overview", make_report())
    assert workbook.sheetnames == [
        "Actual spend",
        "Largest GL accounts",
        "By month",
        "Source",
    ]
    rows = sheet_rows(workbook["By month"])
    assert rows == [
        ["Month", "Actual spend", "Running total", "Lines"],
        ["Total", 0.0, None, 0],
    ]
    assert source_facts(workbook)["Figures as of"] == "No approved uploads yet."


def test_an_export_needs_view_and_export(sample_report, monkeypatch) -> None:
    """Verify a user with view but not export gets 403 from an export address,
    and one with no access 404 (ENG-01 6.2).
    """
    viewer = make_user(fixed_admin=False)
    monkeypatch.setattr(
        access,
        "held",
        lambda who, _: frozenset({access.VIEW}) if who == viewer else frozenset(),
    )
    browser = Client()
    browser.force_login(viewer)
    url = reverse("overview_export", args=[sample_report.pk])
    assert browser.get(url).status_code == 403
    page = browser.get(reverse("overview", args=[sample_report.pk])).content.decode()
    assert "Export to Excel" not in page
    stranger = Client()
    stranger.force_login(make_user(fixed_admin=False))
    assert stranger.get(url).status_code == 404


def test_the_export_link_carries_the_screen_s_filters_keys_and_pick(
    client, sample_report
) -> None:
    """Verify each screen's Export link is its own address plus /export with
    the filters on screen, source rows' keys without the page, and an FAQ
    page's pick written out (ENG-03 9.1, N5).
    """
    base = f"/reports/{sample_report.pk}"
    cases = [
        ("overview", [], "?year=2025", f"{base}/overview/export?year=2025"),
        ("rows", [], "?gl=610100&page=1", f"{base}/rows/export?gl=610100"),
        (
            "faq_page",
            ["gl-account-cost-centers"],
            "?exclude_outside=1",
            f"{base}/faq/gl-account-cost-centers/export?gl=610100&amp;exclude_outside=1",
        ),
    ]
    for name, args, query, link in cases:
        url = reverse(name, args=[sample_report.pk, *args]) + query
        page = client.get(url).content.decode()
        assert f'<a href="{link}">Export to Excel</a>' in page, name


def test_source_rows_export_takes_no_page(client, sample_report) -> None:
    """Verify source rows' export refuses a page, since it holds every row."""
    url = reverse("rows_export", args=[sample_report.pk]) + "?page=1"
    assert client.get(url).status_code == 400


def test_an_export_address_refuses_what_its_screen_refuses(
    client, sample_report
) -> None:
    """Verify an export takes the same values as its screen, refused the
    same way (ENG-03 9.1).
    """
    for name, args, query in [
        ("overview_export", [], "?year=abc"),
        ("gl_account_export", ["0610100"], ""),
        ("faq_page_export", ["total-spend"], "?gl=610100"),
    ]:
        url = reverse(name, args=[sample_report.pk, *args]) + query
        assert client.get(url).status_code == 400, name


def test_the_row_limit_is_excel_s_less_the_header_and_total() -> None:
    """Verify the limit is Excel's 1,048,576 rows, less the header and total
    rows (ENG-03 9.4).
    """
    assert export.MAX_DATA_ROWS == 1_048_576 - 2


def test_the_exceptions_export_says_the_outside_filter_doesn_t_apply(
    client, sample_report
) -> None:
    """Verify the exceptions screen's export says outside IDS isn't filtered
    there, and counts every committed upload's lines with it included, as
    the screen does (ENG-03 8).
    """
    _, workbook = download(
        client, "exceptions", sample_report, query="?exclude_outside=1"
    )
    assert source_facts(workbook)["Filters"] == (
        "Year: all. Outside IDS: included, since this filter doesn't apply here."
    )
    rows = sheet_rows(workbook["Source"])
    (upload,) = [row for row in rows if row and row[0] == 1]
    assert upload[7] == 35


def test_the_source_sheet_names_an_upload_awaiting_approval(
    client, sample_report
) -> None:
    """Verify the Source sheet says an upload awaiting approval isn't included,
    as the as-of line does (ENG-03 3.2).
    """
    make_upload(sample_report, status=Upload.Status.AWAITING_APPROVAL)
    _, workbook = download(client, "overview", sample_report)
    assert source_facts(workbook)["Awaiting approval"] == "Upload 2 isn't included."
