"""Reading a file's bytes into sheets of cells (ENG-02 4.1). No database."""

import io
import zipfile
from collections.abc import Callable
from pathlib import Path

import openpyxl
import pytest
import xlsxwriter
from xlsxwriter.worksheet import Worksheet

from ids_audit.pipeline import readers
from ids_audit.pipeline.formats import TEXT, XLS, XLSX, ZIP_SIGNATURE
from ids_audit.pipeline.readers import (
    FileUnreadable,
    read_calamine,
    read_openpyxl,
    read_sheets,
    read_text,
)

FIXTURES_DIR = Path(__file__).parent / "fixtures"
SAMPLE_XLSX = (FIXTURES_DIR / "sample.xlsx").read_bytes()
SAMPLE_CSV = (FIXTURES_DIR / "sample_semicolon.csv").read_bytes()

# Transactions row 7 (ENG-02 12.3). Year and Period are number cells in the
# workbook, and text in the CSV.
ROW_7 = (
    "1000",
    "0000610100",
    "5100000123",
    "002",
    "KR",
    "28.12.2025",
    "04.01.2026",
    2026,
    1,
    "0000004510",
    "INV-88231",
    "Licence renewal",
    "12,400.000",
    "OMR",
)

BOTH_READERS = [
    pytest.param(read_calamine, id="calamine"),
    pytest.param(read_openpyxl, id="openpyxl"),
]


def make_xlsx(write: Callable[[Worksheet], object]) -> bytes:
    """Return the bytes of a one-sheet workbook XlsxWriter writes with write()."""
    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    write(workbook.add_worksheet("Sheet1"))
    workbook.close()
    return buffer.getvalue()


def rewrite_entry(data: bytes, name: str, old: bytes, new: bytes) -> bytes:
    """Return a copy of an .xlsx, a ZIP file, with old replaced by new in the
    entry called name.
    """
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as copy,
    ):
        for entry in source.infolist():
            body = source.read(entry.filename)
            if entry.filename == name:
                assert body.count(old) == 1
                body = body.replace(old, new)
            copy.writestr(entry, body)
    return buffer.getvalue()


def fail(data: bytes) -> None:
    """Stand in for a reader that can't open the file."""
    raise ValueError("can't read this")


# The sample workbook


def test_calamine_reads_every_sheet_of_the_sample_from_row_1() -> None:
    """Verify the sample's four sheets come in order, and Transactions has 52
    rows of 14 cells, with row 7 at index 6 and numbers as floats.
    """
    sheets = read_calamine(SAMPLE_XLSX)
    names = [sheet.name for sheet in sheets]
    assert names == ["README", "Transactions", "GL Mapping", "Cost Center Mapping"]
    rows = sheets[1].rows
    assert len(rows) == 52
    assert {len(row) for row in rows} == {14}
    assert rows[0][0] == "G/L Account Line Item Display"
    assert rows[2] == ("",) * 14
    assert rows[6] == ROW_7
    assert isinstance(rows[6][7], float)
    assert rows[51][0] == "Total"


def test_openpyxl_reads_the_sample_as_calamine_does() -> None:
    """Verify both readers give equal sheets for the sample, so empty cells are
    "" and short rows are padded in both. 2026 equals 2026.0 in Python.
    """
    assert read_openpyxl(SAMPLE_XLSX) == read_calamine(SAMPLE_XLSX)


@pytest.mark.parametrize("read", BOTH_READERS)
def test_blank_rows_and_columns_at_the_top_keep_their_place(read) -> None:
    """Verify a table starting at B3 is read with rows 1 and 2 and column A
    blank, so its header is row 3 and its first row is row 4 (ENG-02 2).
    """

    def write(sheet: Worksheet) -> None:
        sheet.write_row(2, 1, ["G/L Account", "Cost Center", "Amount"])
        sheet.write_row(3, 1, ["610100", "4510", "5.000"])

    [sheet] = read(make_xlsx(write))
    assert sheet.rows == (
        ("", "", "", ""),
        ("", "", "", ""),
        ("", "G/L Account", "Cost Center", "Amount"),
        ("", "610100", "4510", "5.000"),
    )


def test_formatted_empty_cells_after_the_table_add_nothing() -> None:
    """Verify empty cells with a format, right of the table in its own row and
    far below it, add no columns or rows. calamine ignores them, and openpyxl,
    which reads them, is tidied to match.
    """
    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = workbook.add_worksheet("Sheet1")
    sheet.write_row(0, 0, ["A", "B"])
    # XlsxWriter writes an empty cell only when it has a format.
    bold = workbook.add_format({"bold": True})
    sheet.write_blank(0, 5, None, bold)
    sheet.write_blank(9, 0, None, bold)
    workbook.close()
    data = buffer.getvalue()
    assert read_calamine(data)[0].rows == read_openpyxl(data)[0].rows == (("A", "B"),)


def test_a_sheet_saved_with_the_wrong_size_is_read_in_full() -> None:
    """Verify openpyxl reads every cell even when the file says the sheet is
    only cell A1, as some programs save it.
    """
    wrong = rewrite_entry(
        SAMPLE_XLSX,
        "xl/worksheets/sheet2.xml",
        b'<dimension ref="A1:N52"/>',
        b'<dimension ref="A1"/>',
    )
    assert read_openpyxl(wrong) == read_calamine(SAMPLE_XLSX)


def test_a_chart_sheet_is_left_out() -> None:
    """Verify calamine leaves out a chart sheet, which holds no cells, as
    openpyxl does.
    """
    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    workbook.add_worksheet("Data").write_row(0, 0, [1, 2, 3])
    chart = workbook.add_chart({"type": "column"})
    chart.add_series({"values": "=Data!$A$1:$C$1"})
    workbook.add_chartsheet("Chart").set_chart(chart)
    workbook.close()
    data = buffer.getvalue()
    assert [sheet.name for sheet in read_calamine(data)] == ["Data"]
    assert [sheet.name for sheet in read_openpyxl(data)] == ["Data"]


# Formulas and errors: only saved results are read (ENG-02 4.1). Each test reads
# with both: calamine shows the library's own behaviour, and openpyxl is made to
# match it.


def test_a_formula_reads_as_its_saved_result() -> None:
    """Verify a formula saved with its result gives the result, not its text."""
    data = make_xlsx(lambda sheet: sheet.write_formula(0, 0, "=1+1", None, 2))
    assert read_calamine(data)[0].rows == read_openpyxl(data)[0].rows == ((2,),)


def test_an_error_cell_reads_as_empty() -> None:
    """Verify a #N/A error reads as "" in both readers."""

    def write(sheet: Worksheet) -> None:
        sheet.write_formula(0, 0, "=NA()", None, "#N/A")
        sheet.write_string(0, 1, "x")

    data = make_xlsx(write)
    assert read_calamine(data)[0].rows == read_openpyxl(data)[0].rows == (("", "x"),)


def test_a_formula_without_a_saved_result_reads_as_empty() -> None:
    """Verify a formula openpyxl wrote, which has no saved result, reads as "" in
    both readers.
    """
    workbook = openpyxl.Workbook()
    workbook.active.append(["=1+1", "x"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    assert read_calamine(data)[0].rows == read_openpyxl(data)[0].rows == (("", "x"),)


# Choosing the reader


def test_an_xlsx_calamine_cant_read_falls_back_to_openpyxl(monkeypatch) -> None:
    """Verify an .xlsx calamine fails on is read by openpyxl."""
    monkeypatch.setattr(readers, "read_calamine", fail)
    assert read_sheets(SAMPLE_XLSX, XLSX) == read_openpyxl(SAMPLE_XLSX)


def test_an_xls_calamine_cant_read_is_unreadable(monkeypatch) -> None:
    """Verify an .xls calamine fails on isn't given to openpyxl, which reads
    only .xlsx.
    """
    monkeypatch.setattr(readers, "read_calamine", fail)
    with pytest.raises(FileUnreadable, match="couldn't be opened as Excel"):
        read_sheets(SAMPLE_XLSX, XLS)


def test_a_damaged_xlsx_is_unreadable() -> None:
    """Verify bytes with the ZIP signature that neither reader opens raise
    FileUnreadable, so the upload ends at failed (ENG-02 6).
    """
    with pytest.raises(FileUnreadable, match="couldn't be opened as Excel"):
        read_sheets(ZIP_SIGNATURE + b"not really a zip file", XLSX)


# Delimited text


def test_the_sample_csv_is_one_sheet_of_text() -> None:
    """Verify the sample CSV is one sheet named "", with 52 rows of 14 cells and
    row 7 at index 6, every cell a string.
    """
    [sheet] = read_sheets(SAMPLE_CSV, TEXT)
    assert sheet.name == ""
    assert len(sheet.rows) == 52
    assert {len(row) for row in sheet.rows} == {14}
    assert sheet.rows[6] == tuple(str(cell) for cell in ROW_7)


def test_utf16_text_reads_as_utf8_does() -> None:
    """Verify the sample CSV in UTF-16, with its byte-order mark, gives the same
    sheet.
    """
    utf16 = SAMPLE_CSV.decode("utf-8-sig").encode("utf-16")
    assert read_text(utf16) == read_text(SAMPLE_CSV)


def test_text_that_isnt_utf8_is_unreadable_with_advice() -> None:
    """Verify cp1256 text, with Arabic past the format check's first bytes,
    fails with the message to save it as CSV UTF-8 (ENG-02 4.1).
    """
    lines = ["CoCd;G/L Acct;Text"] + ["1000;0000610100;Support"] * 200
    data = "\r\n".join(lines + ["1000;0000610100;ت"]).encode("cp1256")
    with pytest.raises(FileUnreadable, match="Save it as CSV UTF-8"):
        read_text(data)


def test_a_comma_csv_with_quoted_amounts_at_line_ends() -> None:
    """Verify the file csv.Sniffer gives up on is split by commas, with each
    quoted amount kept whole (ENG-02 4.1).
    """
    data = b"CoCd,G/L Acct,Amount\r\n" + b'1000,610100,"1,234.500"\r\n' * 40
    rows = read_text(data).rows
    assert rows[0] == ("CoCd", "G/L Acct", "Amount")
    assert set(rows[1:]) == {("1000", "610100", "1,234.500")}


def test_a_semicolon_csv_below_bare_title_lines() -> None:
    """Verify title lines without separators, here more than the table's lines,
    don't hide the semicolons below them: one-cell lines never count. This is
    the other file csv.Sniffer gives up on.
    """
    data = b"Title\r\n" * 8 + b"A;B;C\r\n" + b"1;2,5;3\r\n" * 5
    rows = read_text(data).rows
    assert rows[0] == ("Title", "", "")
    assert rows[8] == ("A", "B", "C")
    assert rows[9] == ("1", "2,5", "3")


def test_when_two_separators_split_as_many_lines_more_cells_win() -> None:
    """Verify a tie on lines goes to the separator giving more cells: here the
    comma, which comes after the semicolon in the order tried.
    """
    data = b"A,B,C;D\r\n1,2,3;4\r\n"
    assert read_text(data).rows == (("A", "B", "C;D"), ("1", "2", "3;4"))


def test_a_tab_separated_file() -> None:
    """Verify a tab-separated file is split by tabs."""
    data = b"A\tB\tC\r\n" + b"1\t2,000.000\t3\r\n" * 10
    assert read_text(data).rows[1] == ("1", "2,000.000", "3")


def test_a_line_break_inside_quotes_stays_in_its_row() -> None:
    """Verify a quoted cell holding a line break is one cell, so the next row is
    still row 3, as Excel shows it (ENG-02 2).
    """
    data = b'A;B\r\n1;"two\r\nlines"\r\n3;x\r\n'
    assert read_text(data).rows == (("A", "B"), ("1", "two\r\nlines"), ("3", "x"))


def test_text_no_separator_splits_is_one_column() -> None:
    """Verify a file no separator splits is read as one column."""
    assert read_text(b"Amount\r\n5\r\n").rows == (("Amount",), ("5",))
