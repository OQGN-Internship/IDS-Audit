"""The report screens in a real browser (ENG-03 5 and 11): Playwright drives
Chromium against the application, served by pytest-django's live server.

These run after every other test, as tests/conftest.py arranges, since
Playwright keeps an event loop running in the main thread.
"""

import re
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from playwright.sync_api import Page, expect

from ids_audit.imports.commit import commit_upload
from ids_audit.imports.services import create_report, receive_file
from ids_audit.imports.stage import stage_upload
from ids_audit.ledger.models import GlCiMap
from ids_audit.reports.models import Report
from tests.factories import add_posting, commit, make_admin, make_report, make_upload
from tests.test_browser import SAMPLE, sign_in
from tests.test_queries import JUL

pytestmark = [
    pytest.mark.django_db(transaction=True),
    pytest.mark.usefixtures("uploads_dir"),
]

ROWS_SAID = re.compile(r"(\d+) rows?, adding up to ([-0-9,.]+) OMR\.")


@pytest.fixture(autouse=True, scope="module")
def database_beside_the_browser(django_db_setup):
    """Let the tests use the database while Playwright's event loop runs in
    the same thread, as tests/test_browser.py does.
    """
    with pytest.MonkeyPatch.context() as patch:
        patch.setenv("DJANGO_ALLOW_ASYNC_UNSAFE", "true")
        yield


def sample(admin) -> Report:
    """The sample file, received, read and approved by the administrator."""
    report = create_report("Sample", admin)
    upload = receive_file(report, admin, "sample.xlsx", SAMPLE.read_bytes())
    commit_upload(stage_upload(upload, admin), admin)
    return report


def rows_said(page: Page) -> tuple[str, str]:
    """The count and total the source rows on the page say."""
    found = ROWS_SAID.search(page.locator("p.hint").first.inner_text())
    return found[1], found[2]


def settled(page: Page) -> None:
    """Wait until HTMX has swapped the panel in and settled it, which is when
    the form drawn again inside it answers a change.
    """
    page.wait_for_function(
        "!document.querySelector('.htmx-request, .htmx-swapping, .htmx-settling')"
    )


def click_point(page: Page, chart_id: str, index: int) -> None:
    """Click a chart's bar or point by its place, at the pixel ECharts drew
    it at: the middle of a bar, or a line's point.
    """
    element = page.locator(f'[data-chart="{chart_id}"]')
    element.scroll_into_view_if_needed()
    box = element.bounding_box()
    x, y = page.evaluate(
        """([chartId, index]) => {
            const element = document.querySelector(`[data-chart="${chartId}"]`);
            const chart = echarts.getInstanceByDom(element);
            const points = JSON.parse(document.getElementById(chartId).textContent);
            const value = Number(points[index].value);
            if (element.dataset.kind === "bars") {
                return chart.convertToPixel({seriesIndex: 0}, [value / 2, index]);
            }
            return chart.convertToPixel({seriesIndex: 0}, [index, value]);
        }""",
        [chart_id, index],
    )
    page.mouse.click(box["x"] + x, box["y"] + y)


def test_overview_to_gl_account_to_source_rows(page: Page, live_server) -> None:
    """Verify a user can go from the overview to a GL account, then to the
    rows behind one of its amounts, which add up to it (ENG-03 11).
    """
    admin = make_admin()
    report = sample(admin)
    sign_in(page, live_server, admin.username)
    page.goto(f"{live_server.url}/reports/{report.pk}/")
    expect(page.get_by_role("heading", name="Overview")).to_be_visible()
    largest = page.locator(
        "table", has=page.get_by_role("cell", name="Software licences")
    )
    largest.get_by_role("link", name="610100").click()
    expect(
        page.get_by_role("heading", name="GL account 610100: Software licences")
    ).to_be_visible()
    row = page.locator("tr", has=page.get_by_role("cell", name="4510", exact=True))
    amount = row.locator("td.amount a")
    written = amount.inner_text()
    amount.click()
    expect(page).to_have_url(re.compile(r"/rows\?gl=610100&cc=4510&dept=APP$"))
    assert rows_said(page) == ("2", written)


def test_both_filters_change_the_panel_and_the_export_follows(
    page: Page, live_server, tmp_path: Path
) -> None:
    """Verify the year and outside IDS filters swap the panel and update the
    address, and the Export link then downloads exactly what's on screen:
    its Source sheet states the same filters and its total is the screen's
    (ENG-03 3.1 and 9, N5).
    """
    admin = make_admin()
    report = sample(admin)
    sign_in(page, live_server, admin.username)
    page.goto(f"{live_server.url}/reports/{report.pk}/overview")
    page.select_option("select[name=year]", "2025")
    expect(page.get_by_text("Year: 2025. Outside IDS: included.")).to_be_visible()
    settled(page)
    page.select_option("select[name=exclude_outside]", "1")
    expect(page.get_by_text("Year: 2025. Outside IDS: excluded.")).to_be_visible()
    expect(page).to_have_url(re.compile(r"year=2025&exclude_outside=1$"))
    # The panel swapped in has its charts drawn again.
    expect(page.locator('[data-chart="chart-gl"] canvas')).to_have_count(1)
    # The Actual spend tile, the total first (ENG-03 4).
    shown = page.locator(".tile .figure a").first.inner_text()

    with page.expect_download() as download:
        page.get_by_role("link", name="Export to Excel").click()
    path = tmp_path / download.value.suggested_filename
    download.value.save_as(path)
    workbook = openpyxl.load_workbook(path)
    facts = {
        row[0]: row[1]
        for row in workbook["Source"].iter_rows(values_only=True)
        if row[0]
    }
    assert facts["Filters"] == "Year: 2025. Outside IDS: excluded."
    total = workbook["Actual spend"].cell(
        row=workbook["Actual spend"].max_row, column=2
    )
    assert f"{Decimal(str(total.value)):,.3f}" == shown


def test_a_bar_and_a_running_total_point_open_their_rows(
    page: Page, live_server
) -> None:
    """Verify clicking a GL account's bar opens rows adding up to its amount,
    and clicking a running-total point opens every month up to it, adding up
    to the running total (ENG-03 5).
    """
    admin = make_admin()
    report = sample(admin)
    sign_in(page, live_server, admin.username)
    overview = f"{live_server.url}/reports/{report.pk}/overview"
    page.goto(overview)
    page.wait_for_function("document.querySelector('[data-chart=chart-gl] canvas')")
    # Bars run highest first from the top (ENG-03 5).
    first, second = page.evaluate(
        """() => {
            const chart = echarts.getInstanceByDom(
                document.querySelector('[data-chart=chart-gl]'));
            return [0, 1].map((i) => chart.convertToPixel({yAxisIndex: 0}, i));
        }"""
    )
    assert first < second
    click_point(page, "chart-gl", 0)
    expect(page).to_have_url(re.compile(r"/rows\?gl=610100$"))
    assert rows_said(page)[1] == "1,195,697.539"

    page.goto(overview)
    page.wait_for_function(
        "document.querySelector('[data-chart=chart-running] canvas')"
    )
    click_point(page, "chart-running", 2)
    expect(page).to_have_url(re.compile(r"/rows\?through=2025-09$"))
    assert rows_said(page)[1] == "1,320,955.417"


def test_the_charts_draw_with_the_network_cut_off(page: Page, live_server) -> None:
    """Verify the overview draws its charts with every address outside the
    application refused, since ECharts is kept in the repository (ENG-03 5).
    """
    admin = make_admin()
    report = sample(admin)
    outside = []

    def only_the_application(route) -> None:
        if route.request.url.startswith(live_server.url):
            route.continue_()
        else:
            outside.append(route.request.url)
            route.abort()

    page.route("**/*", only_the_application)
    sign_in(page, live_server, admin.username)
    page.goto(f"{live_server.url}/reports/{report.pk}/overview")
    for chart in ("chart-gl", "chart-running"):
        expect(page.locator(f'[data-chart="{chart}"] canvas')).to_have_count(1)
    assert outside == []


def test_a_negative_gl_account_draws_left_of_zero(page: Page, live_server) -> None:
    """Verify a GL account whose credits exceed its spend draws its bar to the
    left of zero (ENG-03 5).
    """
    admin = make_admin()
    report = make_report()
    upload = make_upload(report)
    add_posting(upload, JUL, "2000.000", gl_code="620200")
    add_posting(upload, JUL, "-3500.000", gl_code="620200")
    add_posting(upload, JUL, "100.000")
    commit(upload, minute=1)
    sign_in(page, live_server, admin.username)
    page.goto(f"{live_server.url}/reports/{report.pk}/overview")
    page.wait_for_function("document.querySelector('[data-chart=chart-gl] canvas')")
    bar, zero, left = page.evaluate(
        """() => {
            const chart = echarts.getInstanceByDom(
                document.querySelector('[data-chart=chart-gl]'));
            const layout = chart.getModel().getSeriesByIndex(0).getData()
                .getItemLayout(1);
            const grid = chart.getModel().getComponent('grid')
                .coordinateSystem.getRect();
            return [layout.x + layout.width,
                chart.convertToPixel({xAxisIndex: 0}, 0), grid.x];
        }"""
    )
    # The bar ends left of zero, and inside the chart, so none of it is cut off.
    assert left <= bar < zero


def test_a_name_with_a_script_shows_as_text_in_the_tooltip(
    page: Page, live_server
) -> None:
    """Verify a GL account named with an image tag shows that text in the
    chart's tooltip, and its script never runs (ENG-03 5, N7).
    """
    admin = make_admin()
    report = make_report()
    upload = make_upload(report)
    name = '<img src=x onerror="window.ran = true">'
    GlCiMap.objects.create(
        upload=upload, report=report, gl_code="610100", gl_name=name, ci_code="2101"
    )
    add_posting(upload, JUL, "1.000")
    commit(upload, minute=1)
    sign_in(page, live_server, admin.username)
    page.goto(f"{live_server.url}/reports/{report.pk}/overview")
    page.wait_for_function("document.querySelector('[data-chart=chart-gl] canvas')")
    # ECharts' own action shows the tooltip as a hover does, through the
    # same formatter.
    page.evaluate(
        """() => echarts.getInstanceByDom(document.querySelector(
            '[data-chart=chart-gl]')).dispatchAction(
                {type: 'showTip', seriesIndex: 0, dataIndex: 0})"""
    )
    expect(page.get_by_text(f"610100 {name}")).to_be_visible()
    assert page.evaluate("window.ran") is None


def test_the_logo_is_drawn_and_a_focused_link_shows_its_ring(
    page: Page, live_server
) -> None:
    """Verify the browser draws OQGN's logo, and the first link reached with
    the keyboard, the logo's, shows the 2px green-text ring 2px outside it
    (ENG-03 4).
    """
    sign_in(page, live_server, make_admin().username)
    logo = page.get_by_role("img", name="OQ Gas Networks")
    expect(logo).to_be_visible()
    assert logo.evaluate("image => image.complete && image.naturalWidth") > 0
    page.keyboard.press("Tab")
    ring = page.evaluate(
        """() => {
            const focused = document.activeElement;
            const style = getComputedStyle(focused);
            return [focused.className, style.outlineStyle, style.outlineWidth,
                    style.outlineColor, style.outlineOffset];
        }"""
    )
    assert ring == ["logo", "solid", "2px", "rgb(88, 125, 57)", "2px"]
